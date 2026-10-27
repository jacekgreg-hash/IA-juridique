"""Import de la base tarifaire officielle TARIC -> données RITA (§10 du cahier des charges).

Source : extraction mensuelle officielle de la base TARIC (Commission européenne,
groupe CIRCABC « TARIC & Quota Data and Information »), fichier eu-taric-*.zip.
La structure des fichiers est documentée dans le PDF CIRCABC « Explanation for
the Taric database extractions » :

* « Nomenclature FR » : hiérarchie des codes marchandises avec désignations FR.
  Le code est « 10 chiffres + suffixe de ligne » ; le suffixe 80 marque les lignes
  de biens réelles, les autres suffixes des lignes intermédiaires structurelles.
  La hiérarchie est donnée par l'indentation (colonne de tirets), en ordre de
  fichier — PAS par préfixe de code (les branches « autres » changent de
  chiffres par rapport à leur parent).
* « Declarable codes » : drapeau IS_LEAF (code déclarable en douane).
* « Duties Import 01-99 » : mesures par code ; le droit de pays tiers est la
  mesure de type 103 (« Third country duty ») à l'origine ERGA OMNES.

Sorties (régénérées, jamais éditées à la main) :
* data/tariff/nomenclature_fr.json — nomenclature FR des codes déclarables
  (CN8 et TARIC10) avec désignations assemblées depuis la hiérarchie et
  droits de pays tiers.
* data/tariff/tariff_base.json — base du moteur tarifaire : les familles et
  discriminants restent définis ci-dessous (logique RITA), mais codes,
  désignations, taux et bases légales proviennent du TARIC officiel.

Usage : python tools/import_taric.py [chemin/vers/eu-taric-YYYY-MM.zip]
Sans argument, prend le eu-taric-*.zip le plus récent dans data/tariff/raw/.
"""
from __future__ import annotations

import json
import re
import sys
import zipfile
import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW_DIR = ROOT / 'data' / 'tariff' / 'raw'
OUT_NOMENCLATURE = ROOT / 'data' / 'tariff' / 'nomenclature_fr.json'
OUT_BASE = ROOT / 'data' / 'tariff' / 'tariff_base.json'

# Droit de pays tiers : type de mesure 103, origine ERGA OMNES (toutes origines
# sans préférence) — taux INDICATIF, le calcul fiscal (§16) dépend de l'origine.
THIRD_COUNTRY_DUTY_TYPE = '103'
ERGA_OMNES = 'ERGA OMNES'

# --- Familles du moteur tarifaire (logique RITA, §9/§11) ----------------------
# Les codes sont les codes TARIC 10 chiffres ; l'importeur refuse de produire la
# base si un code manque dans le TARIC officiel (aucune invention).
FAMILLES = [
    {
        'cle': 'chaussures',
        'natures': ['chaussure de sport', 'chaussure', 'botte'],
        'discriminant': 'tige',
        'question': 'Quelle est la matière de la tige (la partie supérieure) de la chaussure ?',
        'codes': [
            ('6403190000', 'cuir', "tige cuir => 6403 (chaussures de sport, autres que ski)"),
            ('6404110000', 'textile', "dessus textile => 6404 ; l'une des familles les plus taxées du chapitre 64"),
            ('6402190000', 'plastique/caoutchouc', "semelle et dessus en caoutchouc/plastique => 6402"),
        ],
    },
    {
        'cle': 'telephonie',
        'natures': ['téléphone'],
        'discriminant': 'type_appareil',
        'question': "S'agit-il d'un smartphone (écran tactile) ou d'un téléphone classique à clavier ?",
        'codes': [
            ('8517130000', 'smartphone', "smartphone = écran tactile, système d'exploitation complet"),
            ('8517140000', 'classique', 'téléphone classique/feature phone avec clavier physique'),
        ],
    },
    {
        'cle': 'horlogerie',
        'natures': ['montre'],
        'discriminant': 'mouvement',
        'question': "Quel est le mouvement de la montre : automatique/mécanique, à quartz analogique, ou à affichage numérique ?",
        'codes': [
            ('9101210000', 'automatique', "« Automatic » sur le cadran => remontage automatique ; "
                                          'boîte en métaux précieux ou plaqués => 9101'),
            ('9102210000', 'automatique', 'remontage automatique, boîte autre que métaux précieux => 9102'),
            ('9102110000', 'quartz', 'montre fonctionnant électriquement à affichage mécanique '
                                     'seulement (quartz analogique)'),
            ('9102190000', 'numerique', "montre électrique « autres » : comprend l'affichage "
                                        'numérique (LCD) et les autres affichages'),
        ],
    },
    {
        'cle': 'meubles',
        'natures': ['commode', 'armoire', 'étagère', 'canapé', 'fauteuil', 'chaise', 'table', 'lit', 'meuble'],
        'discriminant': 'matiere_principale',
        'question': 'Quelle est la matière principale du meuble : bois, plastique ou métal ?',
        'codes': [
            ('9403609000', 'bois', 'meuble en bois => 9403.6 ; chambres à coucher => 9403 50, '
                                   'salles à manger => 9403 60 10, magasins => 9403 60 30'),
            ('9403700000', 'plastique', 'meuble dont la structure est en plastique'),
            ('9403208080', 'métal', 'meuble à structure métallique => 9403.2 (bureau => 9403.10, '
                                    'lits => 9403 20 20)'),
        ],
    },
]


# --- Utilitaires ----------------------------------------------------------------

def _parse_date(value) -> date | None:
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    s = str(value).strip()
    for fmt in ('%d-%m-%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _aujourdhui(d: date | None) -> bool:
    return d is None or d >= date.today()


def _nettoie_texte(text: str) -> str:
    # '|' marque un retour à la ligne dans la désignation TARIC ; 'n$o' est la
    # graphie d'export de « n° » dans les renvois (ex. « autres que celles du
    # n° 9101 »).
    t = str(text or '').replace('|', ' ').replace('$o', '°')
    return re.sub(r'\s+', ' ', t).strip()


def _profondeur(indent) -> int:
    if not indent:
        return 0
    return len(re.findall(r'-', str(indent)))


def _formate_code(code10: str) -> str:
    """6403190000 -> « 6403 19 00 » (style CN) ; 9403609000 -> « 9403 60 90 ».

    Le TARIC stocke des codes 10 chiffres complétés par des paires de zéros :
    quand les chiffres 9-10 valent 00, le code signifiant est le CN8.
    """
    morceaux = [code10[:4], code10[4:6], code10[6:8]]
    if code10[8:10] != '00':
        morceaux.append(code10[8:10])
    return ' '.join(morceaux)


def _parse_taux(duty: str) -> tuple[float | None, str]:
    """« 16.900 % » -> (0.169, brut) ; « Free » -> (0.0, brut) ; spécifique -> (None, brut)."""
    brut = str(duty or '').strip()
    m = re.search(r'([\d.]+)\s*%', brut)
    if m:
        return round(float(m.group(1)) / 100.0, 4), brut
    if brut.lower() == 'free':
        return 0.0, brut
    return None, brut


def _lire_zip() -> tuple[zipfile.ZipFile, Path]:
    if len(sys.argv) > 1:
        chemin = Path(sys.argv[1])
    else:
        candidats = sorted(RAW_DIR.glob('eu-taric-*.zip'))
        if not candidats:
            sys.exit(f"Aucun eu-taric-*.zip dans {RAW_DIR} — passez le chemin en argument.")
        chemin = candidats[-1]
    print(f"[1/5] Source : {chemin}")
    return zipfile.ZipFile(chemin), chemin


def _version(chemin: Path) -> str:
    m = re.search(r'eu-taric-(\d{4}-\d{2})', chemin.name)
    return m.group(1) if m else chemin.stem


def _reference(chemin: Path) -> str:
    """Date de référence de l'extraction (1er jour du mois publié)."""
    v = _version(chemin)
    if re.fullmatch(r'\d{4}-\d{2}', v):
        return f'{v}-01'
    return date.today().isoformat()


# --- Lecture des tables officielles --------------------------------------------

def _lire_nomenclature(z: zipfile.ZipFile) -> list[dict]:
    """Hiérarchie des codes FR valides, désignations assemblées par indentation.

    Colonne « Hier. Pos. » : niveau du code (2 chapitre, 4 position HS, 6 sous-
    position HS, 8 CN8, 10 TARIC10) ; l'arbre réel suit l'indentation.
    """
    print('[2/5] Lecture de « Nomenclature FR »...')
    wb = load_workbook(io.BytesIO(z.read('Nomenclature FR.xlsx')), read_only=True)
    ws = wb[wb.sheetnames[0]]
    aujourd_hui = date.today()
    # code suffixé -> meilleure ligne valide (description la plus récente)
    lignes: dict[str, dict] = {}
    for code, debut, fin, lang, hier, indent, desc, desc_debut in ws.iter_rows(values_only=True):
        code = str(code or '')
        if not re.fullmatch(r'\d{10} \d{2}', code) or str(lang or '').upper() != 'FR':
            continue
        d_debut = _parse_date(debut)
        d_fin = _parse_date(fin)
        if not d_debut or d_debut > aujourd_hui or not _aujourdhui(d_fin):
            continue  # code non valide aujourd'hui
        cled = code
        cand = {'code10': code[:10], 'suffixe': code[11:], 'niveau': int(hier or 0),
                'profondeur': _profondeur(indent), 'designation': _nettoie_texte(desc),
                'desc_debut': _parse_date(desc_debut)}
        courante = lignes.get(cled)
        if courante is None or (cand['desc_debut'] or date.min) >= (courante['desc_debut'] or date.min):
            lignes[cled] = cand
    wb.close()
    # Ordre du fichier = ordre d'insertion (l'arbre suit l'indentation, cf. doc
    # CIRCABC : le suffixe de ligne « est nécessaire pour construire la structure
    # dans la bonne séquence ») ; le tri code+suffixe la reproduit.
    entrees = sorted(lignes.values(), key=lambda e: (e['code10'], e['suffixe']))
    print(f"      {len(entrees)} lignes valides.")
    return entrees


def _arbre(entrees: list[dict]) -> None:
    """Enrichit chaque entrée : designation_complete + chaîne des codes parents."""
    pile: list[dict] = []
    for e in entrees:
        while pile and pile[-1]['profondeur'] >= e['profondeur']:
            pile.pop()
        parents = list(pile)
        # On assemble depuis la position HS (niveau 4) — les chapitres (niveau 2)
        # portent un intitulé générique, comme dans la désignation CN usuelle.
        fragments = [p['designation'] for p in parents if p['niveau'] >= 4] + [e['designation']]
        e['designation_complete'] = ', '.join(x for x in fragments if x)
        e['codes_parents'] = [p['code10'] for p in parents]
        pile.append(e)


def _lire_declarables(z: zipfile.ZipFile) -> set[str]:
    print('[3/5] Lecture de « Declarable codes »...')
    wb = load_workbook(io.BytesIO(z.read('Declarable codes.xlsx')), read_only=True)
    ws = wb[wb.sheetnames[0]]
    declarables = set()
    for code, _debut, _flag, leaf, fin in ws.iter_rows(values_only=True):
        code = str(code or '')
        if not re.fullmatch(r'\d{10} \d{2}', code):
            continue
        if str(leaf) == '1' and _aujourdhui(_parse_date(fin)):
            declarables.add(code)
    wb.close()
    print(f"      {len(declarables)} codes déclarables.")
    return declarables


def _lire_droits(z: zipfile.ZipFile) -> dict[str, dict]:
    """Droits de pays tiers (103, ERGA OMNES) par code 10 chiffres, valides aujourd'hui."""
    print('[4/5] Lecture de « Duties Import 01-99 » (mesure 103 ERGA OMNES)...')
    wb = load_workbook(io.BytesIO(z.read('Duties Import 01-99.xlsx')), read_only=True)
    ws = wb[wb.sheetnames[0]]
    aujourd_hui = date.today()
    meilleurs: dict[str, dict] = {}
    for code, _add, _order, debut, fin, _red, origine, mtype, legal, duty, _ocode, mcode in ws.iter_rows(values_only=True):
        code = str(code or '')
        if not re.fullmatch(r'\d{10}', code):
            continue
        if str(mcode) != THIRD_COUNTRY_DUTY_TYPE or str(origine) != ERGA_OMNES:
            continue
        d_debut = _parse_date(debut)
        if not d_debut or d_debut > aujourd_hui or not _aujourdhui(_parse_date(fin)):
            continue
        taux, brut = _parse_taux(duty)
        cand = {'taux': taux, 'brut': brut, 'legal': str(legal or '').strip(), 'depuis': d_debut}
        cour = meilleurs.get(code)
        if cour is None or cand['depuis'] >= cour['depuis']:
            meilleurs[code] = cand
    wb.close()
    print(f"      {len(meilleurs)} codes avec droit de pays tiers.")
    return meilleurs


def _droit_cascade(code10: str, parents: list[str], droits: dict[str, dict]) -> dict | None:
    """Droit applicable : sur le code lui-même sinon le premier parent en remontant."""
    for c in [code10] + list(parents)[::-1]:
        if c in droits:
            return droits[c]
    return None


# --- Écriture des sorties -------------------------------------------------------

def _ecrire_nomenclature(entrees, declarables, droits, version, reference):
    print(f"[5/5] Écriture de {OUT_NOMENCLATURE.name} et {OUT_BASE.name}...")
    codes = []
    for e in entrees:
        # Codes déclarables en douane (suffixe 80 + IS_LEAF=1) : CN8 et TARIC10.
        # Le « niveau » hiérarchique TARIC ne suit pas la longueur usuelle : un
        # code CN8 terminé par 00 (ex. 6402 19 00 00) est de niveau 6 (doc
        # CIRCABC §6.2 : niveau = dernière paire de chiffres différente de 00).
        if e['suffixe'] != '80' or f"{e['code10']} 80" not in declarables:
            continue
        droit = _droit_cascade(e['code10'], e['codes_parents'], droits)
        codes.append({
            'code': _formate_code(e['code10']),
            'code10': e['code10'],
            'niveau': e['niveau'],
            'designation': e['designation_complete'],
            'declarable': f"{e['code10']} 80" in declarables,
            'taux_pays_tiers': droit['taux'] if droit else None,
            'droit_pays_tiers_brut': droit['brut'] if droit else None,
            'base_legale': droit['legal'] if droit else None,
        })
    doc = {
        'version': f'TARIC {version}',
        'source': 'Extraction mensuelle officielle de la base TARIC (Commission européenne, '
                  'groupe CIRCABC « TARIC & Quota Data and Information »)',
        'date_reference': reference,
        'langue': 'FR',
        'note': 'Droit de pays tiers = mesure 103 « Third country duty », origine ERGA OMNES ; '
                'taux indicatif, le calcul fiscal dépend de l\'origine (accords, préférences, antidumping).',
        'codes': codes,
    }
    OUT_NOMENCLATURE.write_text(json.dumps(doc, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f"      {len(codes)} codes CN8/TARIC10 -> {OUT_NOMENCLATURE.name} "
          f"({OUT_NOMENCLATURE.stat().st_size / 1e6:.1f} Mo)")
    return {c['code10']: c for c in codes}


def _ecrire_base(nomenclature: dict, version, reference):
    familles = {}
    erreurs = []
    for fam in FAMILLES:
        codes = []
        for code10, valeur, note in fam['codes']:
            e = nomenclature.get(code10)
            if e is None:
                erreurs.append(f"{fam['cle']}: code {code10} absent du TARIC !")
                continue
            entry = {
                'code': e['code'],
                'code10': e['code10'],
                'designation': e['designation'],
                fam['discriminant']: valeur,
                'taux_indicatif_droits': f"{e['taux_pays_tiers']:.3f}"
                if e['taux_pays_tiers'] is not None else None,
                'base_legale': e['base_legale'],
                'notes': note,
            }
            codes.append(entry)
        familles[fam['cle']] = {
            'natures': fam['natures'],
            'discriminant': fam['discriminant'],
            f"question_{fam['discriminant']}": fam['question'],
            'codes': codes,
        }
    if erreurs:
        sys.exit('Codes hors TARIC officiel (import refusé) :\n' + '\n'.join(erreurs))
    base = {
        'version': f'TARIC {version}',
        'source': 'Base TARIC officielle (Commission européenne, extraction CIRCABC '
                  '« TARIC & Quota Data and Information ») — désignations FR et droits de '
                  'pays tiers (mesure 103, ERGA OMNES)',
        'date_validite': reference,
        'familles': familles,
    }
    OUT_BASE.write_text(json.dumps(base, ensure_ascii=False, indent=2), encoding='utf-8')
    print(f"      {len(familles)} familles -> {OUT_BASE.name}")


def main() -> None:
    z, chemin = _lire_zip()
    version, reference = _version(chemin), _reference(chemin)
    entrees = _lire_nomenclature(z)
    _arbre(entrees)
    declarables = _lire_declarables(z)
    droits = _lire_droits(z)
    z.close()
    nomenclature = _ecrire_nomenclature(entrees, declarables, droits, version, reference)
    _ecrire_base(nomenclature, version, reference)
    print('Terminé.')


if __name__ == '__main__':
    main()
